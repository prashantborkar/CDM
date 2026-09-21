import os, re, sys
sys.path.insert(0, os.path.dirname(__file__))
import spec_content as S

HERE = os.path.dirname(os.path.abspath(__file__))
SPEC = os.path.normpath(os.path.join(HERE, ".."))
DIAG = os.path.join(SPEC, "diagrams")
FIGNUM = {k: i + 1 for i, (k, _, _) in enumerate(S.FIGS)}
FIGINFO = {k: (f, c) for k, f, c in S.FIGS}


def sub(t):
    return re.sub(r"\{fig:(\w+)\}", lambda m: "Figure %d" % FIGNUM[m.group(1)], t)


# ------------------------------------------------------------------ summary blocks
def count_reqs():
    fr = {"P1": 0, "P2": 0, "P3": 0}
    fr_pri = {"M": 0, "S": 0, "C": 0}
    ad = nfr = 0
    for b in S.D:
        if b[0] != "tbl":
            continue
        cols, rows = b[1]["cols"], b[1]["rows"]
        if cols[0] != "ID":
            continue
        for r in rows:
            if r[0].startswith("FR-"):
                fr[r[3]] += 1
                fr_pri[r[2]] += 1
            elif r[0].startswith("AD-"):
                ad += 1
            elif r[0].startswith("NFR-"):
                nfr += 1
    return fr, fr_pri, ad, nfr


def summary_blocks():
    fr, fr_pri, ad, nfr = count_reqs()
    total = sum(fr.values())
    B = []
    B.append(("h0", "Summary at a glance"))
    B.append(("p", "The Certificate Deployment Manager (CDM) is a ServiceNow-based platform that keeps every certificate valid, correctly deployed and verified on any "
                   "technology: Windows, Linux, macOS and iOS (through MDM), Kubernetes, Java, load balancers, cloud services and more. Sectigo is the first certificate "
                   "authority; MID Servers in each network zone do the work; a vault holds every credential; and private keys stay on the systems that use them."))
    B.append(("p", "How it works:"))
    B.append(("num", [
        "A certificate is discovered or requested (or a renewal is triggered automatically before expiry).",
        "Policy checks and approvals run. Nothing goes to the CA before that.",
        "The private key and CSR are created on the target. Only the CSR travels.",
        "The CA (Sectigo first) issues the certificate. CDM validates it against the CSR and policy.",
        "The router picks one signed adapter from the target's technology and zone. A MID Server in that zone runs it with a just-in-time credential from the vault.",
        "Pre-check, backup, install, activate, then verify the live endpoint (SAN, chain, expiry, thumbprint). On any failure, roll back automatically.",
        "Inventory, audit trail and notifications are updated. Systems that cannot be automated safely get a manual task, and are still verified automatically.",
    ]))
    B.append(("p", "What the review of your POC folder found (details in chapter 2): the POC proves the routing idea well, and the screenshots give nine solid security and safety principles that are now requirements. "
                   "Eight gaps (G1 to G8) should be fixed before the POC is treated as proof of the design, the most important being: the mock CA returns a fake, non-PEM string; "
                   "no CSR or private key is handled; WinRM runs over HTTP with Basic authentication; one admin password does everything; and script text is pasted into the flow."))
    B.append(("p", "Diagrams (also delivered as separate PNG files in the diagrams folder):"))
    B.append(("tbl", {"cols": ["Figure", "Diagram", "File"],
                      "rows": [[str(i + 1), c, f] for i, (k, f, c) in enumerate(S.FIGS)], "w": [0.7, 5.6, 3.2]}))
    B.append(("p", "Requirement counts:"))
    B.append(("tbl", {"cols": ["Type", "Count", "Detail"], "rows": [
        ["Functional (FR)", str(total), "P1 MVP: %d, P2: %d, P3: %d.  Must: %d, Should: %d, Could: %d." % (fr["P1"], fr["P2"], fr["P3"], fr_pri["M"], fr_pri["S"], fr_pri["C"])],
        ["Platform adapter (AD)", str(ad), "Windows, Linux, Apple, Kubernetes, Java, load balancers, cloud, databases and network."],
        ["Non-functional (NFR)", str(nfr), "Security, availability, performance, audit, operations, usability, portability, crypto-agility."],
        ["Decisions and questions", "15", "Chapter 14, each with a default."]], "w": [2.2, 0.8, 6]}))
    B.append(("note", "warn", "Please review chapter 14 first: fifteen questions with defaults. The most important are Q1 (ServiceNow licences and plugins), Q2 (Sectigo tenant), Q3 (vault) and Q4 (what \"iOS\" means)."))
    return B


BLOCKS = summary_blocks() + S.D

# ------------------------------------------------------------------ numbering
def numbered():
    out = []
    h1n = h2n = 0
    for b in BLOCKS:
        if b[0] == "h1":
            if b[1].startswith("Appendix"):
                out.append(("h1", b[1], b[1]))
            else:
                h1n += 1
                h2n = 0
                out.append(("h1", "%d. %s" % (h1n, b[1]), b[1]))
        elif b[0] == "h2":
            h2n += 1
            out.append(("h2", "%d.%d %s" % (h1n, h2n, b[1]), b[1]))
        else:
            out.append(b)
    return out


NB = numbered()


# ------------------------------------------------------------------ markdown
def slug(t):
    t = t.lower()
    t = re.sub(r"[^\w\s-]", "", t)
    return re.sub(r"\s", "-", t.strip())


def esc(t):
    return sub(t).replace("|", "\\|")


def to_md():
    L = []
    L.append("# %s: %s" % (S.TITLE, S.SUBTITLE))
    L.append("")
    L.append("Version %s | %s | Diagrams are in the [diagrams](diagrams/) folder; a Word version of this document sits next to this file." % (S.VERSION, S.DATE))
    L.append("")
    L.append("## Contents")
    L.append("")
    for b in NB:
        if b[0] == "h0":
            L.append("- [%s](#%s)" % (b[1], slug(b[1])))
        elif b[0] == "h1":
            L.append("- [%s](#%s)" % (b[1], slug(b[1])))
    L.append("")
    for b in NB:
        k = b[0]
        if k == "h0":
            L += ["## " + b[1], ""]
        elif k == "h1":
            L += ["## " + b[1], ""]
        elif k == "h2":
            L += ["### " + b[1], ""]
        elif k == "p":
            L += [sub(b[1]), ""]
        elif k == "bul":
            L += ["- " + sub(i) for i in b[1]] + [""]
        elif k == "num":
            L += ["%d. %s" % (n + 1, sub(i)) for n, i in enumerate(b[1])] + [""]
        elif k == "tbl":
            t = b[1]
            L.append("| " + " | ".join(t["cols"]) + " |")
            L.append("|" + "|".join(["---"] * len(t["cols"])) + "|")
            for r in t["rows"]:
                L.append("| " + " | ".join(esc(c) for c in r) + " |")
            L.append("")
        elif k == "fig":
            f, c = FIGINFO[b[1]]
            n = FIGNUM[b[1]]
            L += ["![Figure %d: %s](diagrams/%s)" % (n, c, f), "", "*Figure %d: %s*" % (n, c), ""]
        elif k == "note":
            L += ["> **Note:** " + sub(b[2]), ""]
        elif k == "code":
            L += ["```json", b[1], "```", ""]
    return "\n".join(L)


# ------------------------------------------------------------------ docx
def to_docx(path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from PIL import Image

    INK = RGBColor(0x1F, 0x2D, 0x3D)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for a in ("left_margin", "right_margin"):
        setattr(sec, a, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    PW = 17.0

    st = doc.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    st["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st["Normal"].paragraph_format.space_after = Pt(6)
    st["Normal"].paragraph_format.line_spacing = 1.1
    for name, size, before, after in (("Heading 1", 18, 18, 8), ("Heading 2", 13.5, 12, 4)):
        s = st[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = INK
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    for name in ("List Bullet", "List Number"):
        st[name].font.name = "Calibri"
        st[name].paragraph_format.space_after = Pt(3)

    def shade(cell, hexcolor):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        tcPr.append(shd)

    def cell_margins(tbl, top=40, bottom=40, left=80, right=80):
        tblPr = tbl._element.tblPr
        m = OxmlElement("w:tblCellMar")
        for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            e = OxmlElement("w:" + side)
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            m.append(e)
        tblPr.append(m)

    def borders(tbl, color="B0BEC5"):
        tblPr = tbl._element.tblPr
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + side)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
            b.append(e)
        tblPr.append(b)

    def set_repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:tblHeader")
        e.set(qn("w:val"), "true")
        trPr.append(e)

    def no_split(row):
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:cantSplit")
        e.set(qn("w:val"), "true")
        trPr.append(e)

    def add_table(t, width_cm=PW, font=8.5):
        cols, rows, w = t["cols"], t["rows"], t["w"]
        tot = float(sum(w))
        widths = [width_cm * x / tot for x in w]
        tb = doc.add_table(rows=1, cols=len(cols))
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tb.autofit = False
        borders(tb)
        cell_margins(tb)
        hdr = tb.rows[0]
        set_repeat_header(hdr)
        for i, c in enumerate(cols):
            cell = hdr.cells[i]
            cell.width = Cm(widths[i])
            shade(cell, "1F2D3D")
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            r = para.add_run(c)
            r.bold = True
            r.font.size = Pt(font)
            r.font.color.rgb = RGBColor(255, 255, 255)
        for ri, row in enumerate(rows):
            rw = tb.add_row()
            no_split(rw)
            for i, c in enumerate(row):
                cell = rw.cells[i]
                cell.width = Cm(widths[i])
                if ri % 2 == 1:
                    shade(cell, "F3F6F8")
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                r = para.add_run(sub(c))
                r.font.size = Pt(font)
                if i == 0 and (c.startswith(("FR-", "AD-", "NFR-", "TC-", "G", "T", "R", "Q", "DD-")) and len(c) < 12):
                    r.bold = True
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_note(kind, text):
        tb = doc.add_table(rows=1, cols=1)
        tb.autofit = False
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(tb, "F2C94C" if kind == "warn" else "90CAF9")
        cell_margins(tb, 80, 80, 140, 140)
        no_split(tb.rows[0])
        cell = tb.rows[0].cells[0]
        cell.width = Cm(PW)
        shade(cell, "FFF8E1" if kind == "warn" else "E8F1FB")
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        r = para.add_run("Note: ")
        r.bold = True
        para.add_run(sub(text))
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_code(text):
        tb = doc.add_table(rows=1, cols=1)
        tb.autofit = False
        borders(tb, "CFD8DC")
        cell_margins(tb, 60, 60, 120, 120)
        cell = tb.rows[0].cells[0]
        cell.width = Cm(PW)
        shade(cell, "F5F7F8")
        first = True
        for line in text.split("\n"):
            para = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            r = para.add_run(line if line else " ")
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(7.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_number_footer(section):
        f = section.footer
        para = f.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run("%s: %s | v%s | Page " % (S.TITLE, S.SUBTITLE, S.VERSION.split(" ")[0]))
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x60, 0x6B, 0x7A)
        r2 = para.add_run()
        for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if typ:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), typ)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
            r2._r.append(e)
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor(0x60, 0x6B, 0x7A)

    page_number_footer(sec)

    # ---------------- title page
    for _ in range(6):
        doc.add_paragraph()
    p0 = doc.add_paragraph()
    r = p0.add_run(S.TITLE)
    r.font.size = Pt(30)
    r.bold = True
    r.font.color.rgb = INK
    p1 = doc.add_paragraph()
    r = p1.add_run(S.SUBTITLE)
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
    p2 = doc.add_paragraph()
    r = p2.add_run("Generic, secure certificate lifecycle automation for Windows, Linux, Apple, Kubernetes and more, on ServiceNow, MID Server and Sectigo")
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(0x60, 0x6B, 0x7A)
    doc.add_paragraph()
    add_table({"cols": ["Item", "Detail"], "rows": [
        ["Version", S.VERSION], ["Date", S.DATE], ["Status", "Draft for review; decisions needed in chapter 14"],
        ["Based on", "Promp.md.txt, POC_SETUP_STATUS.md, mock_sectigo.ps1, data1.jpeg, data2.jpeg, data3.jpeg"],
        ["Companion files", "CDM_Requirements_Specification.md (same content), diagrams folder (11 PNG images)"]], "w": [1.5, 5]}, width_cm=12, font=9.5)
    doc.add_page_break()

    # ---------------- contents
    hh = doc.add_paragraph()
    r = hh.add_run("Contents")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = INK
    for b in NB:
        if b[0] in ("h0", "h1"):
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.line_spacing = 1.0
            rr0 = para.add_run(b[1])
            rr0.font.size = Pt(10)
            rr0.bold = True
        elif b[0] == "h2":
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.8)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            rr = para.add_run(b[1])
            rr.font.size = Pt(8)
            rr.font.color.rgb = RGBColor(0x60, 0x6B, 0x7A)
    doc.add_page_break()

    # ---------------- body
    landscape = False

    def new_section(land):
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        if land:
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = Cm(29.7), Cm(21.0)
            s.left_margin = s.right_margin = Cm(1.5)
            s.top_margin, s.bottom_margin = Cm(1.3), Cm(1.5)
        else:
            s.orientation = WD_ORIENT.PORTRAIT
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
            s.left_margin = s.right_margin = Cm(2.0)
            s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.8)
        return s

    first_h1 = True
    for b in NB:
        k = b[0]
        if k != "fig" and landscape:
            new_section(False)
            landscape = False
        if k == "h0":
            doc.add_heading(b[1], level=1)
        elif k == "h1":
            if not first_h1 and not b[1].startswith("Appendix"):
                pass
            first_h1 = False
            h = doc.add_heading(b[1], level=1)
            pPr = h._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bt = OxmlElement("w:bottom")
            bt.set(qn("w:val"), "single")
            bt.set(qn("w:sz"), "8")
            bt.set(qn("w:space"), "2")
            bt.set(qn("w:color"), "2E7D32")
            bd.append(bt)
            pPr.append(bd)
        elif k == "h2":
            doc.add_heading(b[1], level=2)
        elif k == "p":
            doc.add_paragraph(sub(b[1]))
        elif k == "bul":
            for i in b[1]:
                doc.add_paragraph(sub(i), style="List Bullet")
        elif k == "num":
            for n, i in enumerate(b[1]):
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(0.9)
                para.paragraph_format.first_line_indent = Cm(-0.6)
                para.paragraph_format.space_after = Pt(3)
                rr = para.add_run("%d.  " % (n + 1))
                rr.bold = True
                para.add_run(sub(i))
        elif k == "tbl":
            add_table(b[1])
        elif k == "note":
            add_note(b[1], b[2])
        elif k == "code":
            add_code(b[1])
        elif k == "fig":
            f, c = FIGINFO[b[1]]
            n = FIGNUM[b[1]]
            if not landscape:
                new_section(True)
                landscape = True
            else:
                doc.add_page_break()
            im = Image.open(os.path.join(DIAG, f))
            w_px, h_px = im.size
            max_w, max_h = 26.6, 17.4
            wcm = max_w
            hcm = wcm * h_px / w_px
            if hcm > max_h:
                hcm = max_h
                wcm = hcm * w_px / h_px
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(2)
            para.add_run().add_picture(os.path.join(DIAG, f), width=Cm(wcm))
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = cp.add_run("Figure %d: %s" % (n, c))
            rr.italic = True
            rr.font.size = Pt(9)
            rr.font.color.rgb = RGBColor(0x60, 0x6B, 0x7A)
    doc.core_properties.title = "%s: %s" % (S.TITLE, S.SUBTITLE)
    doc.core_properties.subject = "Requirements specification"
    doc.save(path)


if __name__ == "__main__":
    md = to_md()
    with open(os.path.join(SPEC, "CDM_Requirements_Specification.md"), "w", encoding="utf-8") as fh:
        fh.write(md)
    print("wrote md", len(md))
    to_docx(os.path.join(SPEC, "CDM_Requirements_Specification.docx"))
    print("wrote docx")
